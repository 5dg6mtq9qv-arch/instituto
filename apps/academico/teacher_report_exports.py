"""Exportaciones presentables del informe mensual docente."""

from io import BytesIO

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Image,
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


BLUE = "2563EB"
DARK = "172033"
MUTED = "667085"
LIGHT = "EEF4FF"
GREEN = "059669"
RED = "DC2626"


def metric_value(value, suffix="%"):
    return "-" if value is None else f"{value:.1f}{suffix}"


def teacher_name(teacher):
    return teacher.nombre_completo()


def export_teacher_report_excel(report, institution_name, month_label):
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Resumen docente"
    detail = workbook.create_sheet("Detalle por materia")
    _build_summary_sheet(summary, report, institution_name, month_label)
    _build_detail_sheet(detail, report, institution_name, month_label)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = (
        f'attachment; filename="informe_docentes_{report["start"]:%Y_%m}.xlsx"'
    )
    return response


def _title_rows(sheet, institution_name, month_label, column_count):
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=column_count)
    title = sheet.cell(1, 1, institution_name)
    title.font = Font(size=16, bold=True, color="FFFFFF")
    title.fill = PatternFill("solid", fgColor=DARK)
    title.alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[1].height = 28
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=column_count)
    subtitle = sheet.cell(2, 1, f"Informe mensual docente · {month_label}")
    subtitle.font = Font(size=12, bold=True, color=DARK)
    subtitle.fill = PatternFill("solid", fgColor=LIGHT)
    subtitle.alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[2].height = 23


def _header(sheet, row, headers):
    for column, label in enumerate(headers, start=1):
        cell = sheet.cell(row=row, column=column, value=label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[row].height = 34


def _body_style(sheet, start_row, end_row, column_count):
    border = Border(bottom=Side(style="thin", color="D7DEE8"))
    for row in range(start_row, end_row + 1):
        if row % 2 == 0:
            for column in range(1, column_count + 1):
                sheet.cell(row, column).fill = PatternFill("solid", fgColor="F8FAFC")
        for column in range(1, column_count + 1):
            cell = sheet.cell(row, column)
            cell.border = border
            cell.alignment = Alignment(
                horizontal="left" if column <= 3 else "center",
                vertical="center",
                wrap_text=True,
            )


def _percentage_cell(sheet, row, column, value):
    cell = sheet.cell(row, column, None if value is None else value / 100)
    cell.number_format = "0.0%"


def _build_summary_sheet(sheet, report, institution_name, month_label):
    headers = [
        "Docente", "Clases", "Asistencia docente", "Faltas", "Reemplazadas",
        "Pendientes", "Aula virtual", "Asistencia estudiantes", "Promedio Moodle",
        "Avance temario", "Aprobadas", "En revisión", "Rechazadas", "Atrasadas",
    ]
    _title_rows(sheet, institution_name, month_label, len(headers))
    _header(sheet, 4, headers)
    for row_number, row in enumerate(report["rows"], start=5):
        metrics = row["metrics"]
        values = [
            teacher_name(row["teacher"]), metrics["scheduled"], None, metrics["absent"],
            metrics["replaced"], metrics["teacher_pending"], None, None, None, None,
            metrics["plan"]["aprobada"], metrics["plan"]["revision"],
            metrics["plan"]["rechazada"], metrics["plan_late"],
        ]
        for column, value in enumerate(values, start=1):
            sheet.cell(row_number, column, value)
        _percentage_cell(sheet, row_number, 3, metrics["teacher_attendance_percent"])
        _percentage_cell(sheet, row_number, 7, metrics["virtual_percent"])
        _percentage_cell(sheet, row_number, 8, metrics["student_attendance_percent"])
        _percentage_cell(sheet, row_number, 9, metrics["grade_average"])
        _percentage_cell(sheet, row_number, 10, metrics["topic_percent"])

    end_row = max(5, 4 + len(report["rows"]))
    _body_style(sheet, 5, end_row, len(headers))
    sheet.freeze_panes = "A5"
    sheet.auto_filter.ref = f"A4:{get_column_letter(len(headers))}{end_row}"
    widths = [34, 10, 17, 9, 13, 12, 14, 19, 16, 15, 11, 11, 12, 11]
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.sheet_view.showGridLines = False
    sheet.print_title_rows = "1:4"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True


def _build_detail_sheet(sheet, report, institution_name, month_label):
    headers = [
        "Docente", "Grupo", "Materia", "Clases", "Asist. docente", "Aula virtual",
        "Asist. estudiantes", "Promedio", "Cobertura notas", "Temario",
        "Subtemas del mes", "Aprobadas", "Rechazadas", "Atrasadas",
    ]
    _title_rows(sheet, institution_name, month_label, len(headers))
    _header(sheet, 4, headers)
    current_row = 5
    for row in report["rows"]:
        for metrics in row["subjects"]:
            values = [
                teacher_name(row["teacher"]), metrics["grupo"].nombre, metrics["materia"].nombre,
                metrics["scheduled"], None, None, None, None, None, None,
                metrics["monthly_subtopics_count"], metrics["plan"]["aprobada"],
                metrics["plan"]["rechazada"], metrics["plan_late"],
            ]
            for column, value in enumerate(values, start=1):
                sheet.cell(current_row, column, value)
            for column, key in [
                (5, "teacher_attendance_percent"), (6, "virtual_percent"),
                (7, "student_attendance_percent"), (8, "grade_average"),
                (9, "grade_coverage_percent"), (10, "topic_percent"),
            ]:
                _percentage_cell(sheet, current_row, column, metrics[key])
            current_row += 1
    end_row = max(5, current_row - 1)
    _body_style(sheet, 5, end_row, len(headers))
    sheet.freeze_panes = "A5"
    sheet.auto_filter.ref = f"A4:{get_column_letter(len(headers))}{end_row}"
    widths = [32, 25, 25, 9, 15, 14, 18, 12, 16, 12, 15, 11, 11, 11]
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True


def export_teacher_report_pdf(report, institution, month_label):
    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=12 * mm,
        leftMargin=12 * mm,
        topMargin=13 * mm,
        bottomMargin=14 * mm,
        title=f"Informe mensual docente - {month_label}",
        author=institution.nombre_display() if institution else "Instituto",
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ReportTitle", parent=styles["Title"], fontName="Helvetica-Bold",
        fontSize=19, leading=22, textColor=colors.HexColor(f"#{DARK}"), alignment=TA_LEFT,
    )
    subtitle_style = ParagraphStyle(
        "ReportSubtitle", parent=styles["Normal"], fontSize=10, leading=14,
        textColor=colors.HexColor(f"#{MUTED}"), alignment=TA_LEFT,
    )
    section_style = ParagraphStyle(
        "ReportSection", parent=styles["Heading2"], fontName="Helvetica-Bold",
        fontSize=13, leading=16, textColor=colors.HexColor(f"#{DARK}"), spaceAfter=7,
    )
    small_style = ParagraphStyle(
        "Small", parent=styles["Normal"], fontSize=7.2, leading=9,
        textColor=colors.HexColor(f"#{DARK}"), alignment=TA_CENTER,
    )
    small_left = ParagraphStyle("SmallLeft", parent=small_style, alignment=TA_LEFT)

    story = []
    brand_name = institution.nombre_display() if institution else "Instituto"
    header_items = []
    if institution and institution.logo:
        try:
            header_items.append(Image(institution.logo.path, width=22 * mm, height=22 * mm))
        except (OSError, ValueError):
            pass
    header_copy = [
        Paragraph(brand_name, title_style),
        Paragraph(f"Informe mensual docente · {month_label}", subtitle_style),
        Paragraph(
            f"Corte de información: {report['cutoff']:%d/%m/%Y}. "
            "Las clases futuras no se incluyen.", subtitle_style,
        ),
    ]
    header_items.append(header_copy)
    header = Table([header_items], colWidths=[26 * mm, 235 * mm] if len(header_items) == 2 else [261 * mm])
    header.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(f"#{LIGHT}")),
        ("BOX", (0, 0), (-1, -1), 0.6, colors.HexColor("#CBD5E1")),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
    ]))
    story.extend([header, Spacer(1, 7 * mm)])

    summary = report["summary"]
    cards = [
        ("Docentes", str(summary["teachers"])),
        ("Clases", str(summary["scheduled"])),
        ("Asistencia docente", metric_value(summary["teacher_attendance_percent"])),
        ("Aula virtual", metric_value(summary["virtual_percent"])),
        ("Asistencia alumnos", metric_value(summary["student_attendance_percent"])),
        ("Planificaciones tarde", str(summary["late_plans"])),
    ]
    card_table = Table(
        [[Paragraph(label, small_style), Paragraph(f"<b>{value}</b>", subtitle_style)] for label, value in cards],
        colWidths=[43.5 * mm] * 6,
    )
    card_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.white),
        ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#D7DEE8")),
        ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D7DEE8")),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    story.extend([card_table, Spacer(1, 6 * mm), Paragraph("Resumen comparativo", section_style)])
    story.append(_pdf_summary_table(report, small_style, small_left))

    if len(report["rows"]) == 1:
        story.extend([PageBreak(), Paragraph(
            f"Detalle de {teacher_name(report['rows'][0]['teacher'])}", section_style
        )])
        story.append(_pdf_detail_table(report["rows"][0], small_style, small_left))
        story.append(Spacer(1, 6 * mm))
        story.append(KeepTogether([
            Paragraph("Criterios de lectura", section_style),
            Paragraph(
                "La asistencia estudiantil usa listas cerradas y excluye justificativos del denominador. "
                "Las notas corresponden a actividades calificadas en el mes y se normalizan sobre 100. "
                "El avance del temario es acumulado hasta la fecha de corte.", subtitle_style,
            ),
        ]))

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor("#D7DEE8"))
        canvas.line(12 * mm, 10 * mm, 285 * mm, 10 * mm)
        canvas.setFillColor(colors.HexColor(f"#{MUTED}"))
        canvas.setFont("Helvetica", 7)
        canvas.drawString(12 * mm, 6 * mm, "Documento generado por el sistema académico")
        canvas.drawRightString(285 * mm, 6 * mm, f"Página {doc.page}")
        canvas.restoreState()

    document.build(story, onFirstPage=footer, onLaterPages=footer)
    output.seek(0)
    response = HttpResponse(output.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = (
        f'attachment; filename="informe_docentes_{report["start"]:%Y_%m}.pdf"'
    )
    return response


def _pdf_summary_table(report, style, left_style):
    headers = [
        "Docente", "Clases", "Asist. doc.", "Faltas", "Reempl.", "Aula virtual",
        "Asist. alumnos", "Promedio", "Temario", "Aprob.", "Rech.", "Tarde",
    ]
    data = [[Paragraph(f"<b>{header}</b>", style) for header in headers]]
    for row in report["rows"]:
        metrics = row["metrics"]
        data.append([
            Paragraph(teacher_name(row["teacher"]), left_style),
            metrics["scheduled"], metric_value(metrics["teacher_attendance_percent"]),
            metrics["absent"], metrics["replaced"], metric_value(metrics["virtual_percent"]),
            metric_value(metrics["student_attendance_percent"]), metric_value(metrics["grade_average"]),
            metric_value(metrics["topic_percent"]), metrics["plan"]["aprobada"],
            metrics["plan"]["rechazada"], metrics["plan_late"],
        ])
    table = Table(data, repeatRows=1, colWidths=[46 * mm] + [19.5 * mm] * 11)
    table.setStyle(_pdf_table_style())
    return table


def _pdf_detail_table(row, style, left_style):
    headers = [
        "Grupo / materia", "Clases", "Asist. doc.", "Aula virtual", "Asist. alumnos",
        "Promedio", "Cobertura notas", "Temario", "Subtemas mes", "Aprob.", "Rech.", "Tarde",
    ]
    data = [[Paragraph(f"<b>{header}</b>", style) for header in headers]]
    for metrics in row["subjects"]:
        data.append([
            Paragraph(f"<b>{metrics['materia'].nombre}</b><br/>{metrics['grupo'].nombre}", left_style),
            metrics["scheduled"], metric_value(metrics["teacher_attendance_percent"]),
            metric_value(metrics["virtual_percent"]), metric_value(metrics["student_attendance_percent"]),
            metric_value(metrics["grade_average"]), metric_value(metrics["grade_coverage_percent"]),
            metric_value(metrics["topic_percent"]), metrics["monthly_subtopics_count"],
            metrics["plan"]["aprobada"], metrics["plan"]["rechazada"], metrics["plan_late"],
        ])
    table = Table(data, repeatRows=1, colWidths=[48 * mm] + [19.3 * mm] * 11)
    table.setStyle(_pdf_table_style())
    return table


def _pdf_table_style():
    return TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{BLUE}")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ALIGN", (1, 1), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D7DEE8")),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
    ])
